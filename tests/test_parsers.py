from docx import Document
from pptx import Presentation
from pptx.util import Inches

from src.ingestion.parsers import parse_docx, parse_file, parse_pptx


def test_docx_tables_are_extracted_in_document_order(tmp_path):
    doc = Document()
    doc.add_paragraph("before the table")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "replicas"
    table.cell(1, 1).text = "desired Pod count"
    doc.add_paragraph("after the table")
    path = tmp_path / "t.docx"
    doc.save(path)
    lines = parse_docx(path).splitlines()
    assert lines == ["before the table", "Field | Meaning", "replicas | desired Pod count", "after the table"]


def test_pptx_tables_and_text_are_extracted(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Pods"
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    shape.table.cell(0, 0).text = "Kind"
    shape.table.cell(0, 1).text = "Use"
    shape.table.cell(1, 0).text = "Pod"
    shape.table.cell(1, 1).text = "smallest unit"
    path = tmp_path / "t.pptx"
    presentation.save(path)
    text = parse_pptx(path)
    assert "Pods" in text
    assert "Kind | Use" in text
    assert "Pod | smallest unit" in text


def test_parse_file_rejects_unknown_extensions(tmp_path):
    path = tmp_path / "x.bin"
    path.write_text("x")
    try:
        parse_file(path)
    except ValueError:
        return
    raise AssertionError("expected ValueError")
