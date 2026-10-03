from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx import Presentation
from bs4 import BeautifulSoup


def parse_pdf(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n\n".join(page.extract_text() or "" for page in pdf.pages)


def _table_rows(table) -> list[str]:
    return [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]


def parse_docx(path: Path) -> str:
    doc = DocxDocument(path)
    parts: list[str] = []
    for block in doc.iter_inner_content():
        if isinstance(block, DocxParagraph):
            parts.append(block.text)
        else:
            parts.extend(_table_rows(block))
    return "\n".join(parts)


def parse_pptx(path: Path) -> str:
    presentation = Presentation(path)
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.extend(_table_rows(shape.table))
    return "\n".join(parts)


def parse_html(path: Path) -> str:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    return soup.get_text(separator="\n")


def parse_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".html": parse_html,
    ".htm": parse_html,
    ".txt": parse_txt,
    ".md": parse_txt,
    ".yaml": parse_txt,
    ".yml": parse_txt,
}


def parse_file(path: Path) -> str:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"no parser registered for extension: {path.suffix}")
    return parser(path)
